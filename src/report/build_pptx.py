"""스마트 안전장비 도입 효과 분석 발표자료(.pptx) 생성.

visualizations/ 의 차트 PNG와 docu/비율검정_분석문서.md 의 내용을 슬라이드로 조립한다.
결과: report/스마트안전장비_도입효과_분석.pptx

출처: 류정·박인선(2025), Crisisonomy 21(6) / News1(2024)
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE_DIR = Path(__file__).resolve().parent.parent.parent
VIZ = BASE_DIR / "visualizations"
OUT_DIR = BASE_DIR / "report"
OUT_DIR.mkdir(exist_ok=True)

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0xC4, 0x4E, 0x52)

# 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

BLANK = prs.slide_layouts[6]


def _set(run, size, bold=False, color=NAVY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT):
    """lines: [(text, size, bold, color), ...] — 한 문단씩."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = text
        _set(run, size, bold, color)
        p.space_after = Pt(6)
    return tb


def add_title(slide, text):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), SW, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.5)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    _set(run, 28, True, RGBColor(0xFF, 0xFF, 0xFF))


def add_image_slide(title, img_path, caption=None):
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, title)
    if not img_path.exists():
        add_textbox(slide, Inches(1), Inches(3), Inches(11), Inches(1),
                    [(f"[이미지 없음: {img_path.name}]", 18, False, ACCENT)])
        return slide
    # 제목 아래 영역에 이미지 맞춤 (16:9, 좌우 여백)
    from PIL import Image
    with Image.open(img_path) as im:
        iw, ih = im.size
    max_w, max_h = Inches(11.5), Inches(5.3)
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    left = int((SW - w) / 2)
    top = Inches(1.35)
    slide.shapes.add_picture(str(img_path), left, top, width=w, height=h)
    if caption:
        add_textbox(slide, Inches(0.6), Inches(6.9), Inches(12), Inches(0.5),
                    [(caption, 11, False, GRAY)])
    return slide


def add_bullets_slide(title, bullets, subtitle=None):
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, title)
    top = Inches(1.5)
    if subtitle:
        add_textbox(slide, Inches(0.7), top, Inches(12), Inches(0.6),
                    [(subtitle, 18, True, ACCENT)])
        top = Inches(2.2)
    lines = [(f"•  {b}", 18, False, RGBColor(0x22, 0x22, 0x22)) for b in bullets]
    add_textbox(slide, Inches(0.9), top, Inches(11.7), Inches(5), lines)
    return slide


def build():
    # 1. 표지
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, Inches(0), Inches(0), SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    add_textbox(s, Inches(1), Inches(2.6), Inches(11.3), Inches(2),
                [("스마트 안전장비 도입 효과 분석", 40, True, RGBColor(0xFF, 0xFF, 0xFF)),
                 ("비율검정(Two-proportion test) 기반 통계 검증", 22, False, RGBColor(0xCF, 0xDA, 0xE8))])
    add_textbox(s, Inches(1), Inches(6.2), Inches(11.3), Inches(0.8),
                [("K-NAVI · industrial-data  |  근거: 류정·박인선(2025), Crisisonomy 21(6) / News1(2024)",
                  14, False, RGBColor(0xAF, 0xBE, 0xD0))])

    # 2. 배경·목표
    add_bullets_slide(
        "배경 · 목표",
        ["가설: 스마트장비·로봇 도입이 산업재해(재해율)를 줄이는가?",
         "접근: 도입 현장과 전국 평균의 재해율을 정형데이터 분석으로 비교",
         "재해율은 '재해 발생 여부'의 비율 데이터 → 평균 비교(t-검정)가 아니라 비율검정이 적합",
         "결과를 K-NAVI(정밀 인지·소통 보완 기술)의 필요성 근거로 연결"],
        subtitle="로봇 도입만으로 사고가 줄어드는지 통계로 확인한다")

    # 3. 근거 데이터
    add_bullets_slide(
        "근거 데이터",
        ["전국 중소규모 건설현장(2023): 근로자 576,224명 · 사고 2,656명(사망 171 + 부상 2,485) · 재해율 0.461%",
         "스마트 안전장비 지원현장 123개소(2021~2023): 근로자 4,481명 · 사고 16명(사망 1 + 부상 15) · 재해율 0.357%",
         "출처: 류정·박인선(2025), 「건설현장의 스마트 안전장비 성능과 도입 효과 분석」, Crisisonomy / News1(2024)",
         "배경 통계: 고용노동부 산업재해현황, IFR 로봇밀도"])

    # 4~5. 도입 효과 차트
    add_image_slide("도입 효과 (1) — 사고재해율 비교",
                    VIZ / "10_스마트안전장비_지원현장_사고재해율_비교.png",
                    "전국 평균 0.461% → 지원현장 0.357% (상대 -22.56%) → 본 연구 현장 0%")
    add_image_slide("도입 효과 (2) — 사망자 · 부상자 수",
                    VIZ / "11_스마트안전장비_사망자_부상자_비교.png",
                    "절대 건수는 모집단 규모(576,224 vs 4,481) 차이를 반영 — 규모 보정 비교는 재해율(%) 기준")

    # 6. 검증 방법
    add_bullets_slide(
        "검증 방법 — 비율검정",
        ["귀무가설 H0: 전국과 도입현장의 재해율이 같다 (효과 없음)",
         "대립가설 H1: 도입현장 재해율이 더 낮다 (단측, α=0.05)",
         "① 2표본 비율 Z-검정 (정규근사)",
         "② Fisher 정확검정 (사고가 드문 사건이라 근사 없이 정확 확률로 재확인)",
         "③ 95% 신뢰구간으로 차이의 크기·불확실성 제시"],
        subtitle="H0는 목표와 반대로 '효과 없음'에 두고, 데이터로 뒤집을 수 있는지 본다")

    # 7~8. 검정 시각화
    add_image_slide("검증 (1) — 2표본 비율 Z-검정",
                    VIZ / "proportion_test_R" / "01_Z검정_정규분포_기각역.png",
                    "관측 z=1.023 < 임계값 1.645 → 기각역 밖 → 유의하지 않음 (단측 p=0.153)")
    add_image_slide("검증 (2) — Fisher 정확검정",
                    VIZ / "proportion_test_R" / "02_Fisher_초기하분포_p값.png",
                    "관측 16건 vs H0 기대 20.6건 · P(X≤16)=0.182 = Fisher p · OR=0.774")

    # 9. 종합 결과
    slide = prs.slides.add_slide(BLANK)
    add_title(slide, "종합 결과")
    rows = [
        ("검정 방법", "통계량", "유의 여부(α=0.05)"),
        ("2표본 비율 Z-검정 (단측)", "z=1.023, p=0.153", "유의하지 않음"),
        ("Fisher 정확검정 (단측)", "OR=0.774, p=0.182", "유의하지 않음"),
        ("95% 신뢰구간 (차이)", "[-0.072%p, 0.279%p]", "0 포함 → 유의하지 않음"),
    ]
    table = slide.shapes.add_table(len(rows), 3, Inches(1.2), Inches(1.9),
                                   Inches(10.9), Inches(3)).table
    table.columns[0].width = Inches(4.3)
    table.columns[1].width = Inches(3.6)
    table.columns[2].width = Inches(3.0)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            _set(run, 15, bold=(r == 0),
                 color=RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else RGBColor(0x22, 0x22, 0x22))
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else RGBColor(0xF2, 0xF4, 0xF7)
    add_textbox(slide, Inches(1.2), Inches(5.3), Inches(11), Inches(1),
                [("세 방법 모두 결론 일치: 관찰된 감소는 방향성만 있고 통계적으로 유의하지 않음",
                  18, True, ACCENT)])

    # 10. 결론
    add_bullets_slide(
        "결론 · 해석",
        ["도입현장 재해율이 전국 평균 대비 22.56% 낮게 관찰됨 (0.357% vs 0.461%)",
         "그러나 통계적 유의성은 확인되지 않음 (Fisher p≈0.18, 95% CI가 0 포함)",
         "원인: 사고는 드문 사건 + 도입현장 표본(4,481명·16건)이 작아 검정력 부족",
         "'유의하지 않음'은 '효과 없음'의 증명이 아니라 '판단 보류'",
         "→ 방향성 있는 정황 근거로서 K-NAVI 같은 정밀 안전기술과 데이터 축적의 필요성을 뒷받침"],
        subtitle="방향은 감소, 그러나 확정 입증에는 데이터가 더 필요")

    # 11. 한계
    add_bullets_slide(
        "한계",
        ["도입현장 123개소는 무작위 표본이 아닌 지원사업 선정 현장 → 선택 편향 가능성",
         "교란변수(현장 규모·관리 수준 등) 통제 없음 → 인과 단정 불가",
         "전국 근로자 수 576,224명은 논문 인용 기사 원자료 기준(본문 '약 570만'은 표기 오류)",
         "표본·사건 수 확대 시 검정력 개선 여지"])

    out = OUT_DIR / "스마트안전장비_도입효과_분석.pptx"
    prs.save(str(out))
    print(f"완료: {out}  (슬라이드 {len(prs.slides._sldIdLst)}장)")


if __name__ == "__main__":
    build()
